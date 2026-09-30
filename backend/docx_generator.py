"""Markdown-ish text -> DOCX (python-docx)."""
import re

from docx import Document
from docx.shared import Pt, RGBColor


def _runs(paragraph, text):
    for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
        elif part:
            paragraph.add_run(part)


def build_docx(title, body, path):
    doc = Document()
    doc.add_heading(title, level=0)
    in_code, buf = False, []
    for line in body.splitlines():
        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(buf))
                run.font.name, run.font.size = "Consolas", Pt(9)
                run.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
                buf = []
            in_code = not in_code
            continue
        if in_code:
            buf.append(line)
        elif line.startswith("#"):
            doc.add_heading(line.lstrip("# ").strip(), level=min(line.count("#", 0, 4), 3))
        elif re.match(r"^\s*[-*] ", line):
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\s*\d+[.)] ", line):
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)] ", "", line))
        elif line.strip():
            _runs(doc.add_paragraph(), line)
    doc.save(str(path))
